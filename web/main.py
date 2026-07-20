#!/usr/bin/env python3
"""G.A.B.I. Web — FastAPI backend."""

import contextlib
import functools
import json
import os
import re as _re
import secrets
import sqlite3
import tempfile
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Annotated
from zoneinfo import ZoneInfo

import aiofiles
import anyio
import bcrypt
import openpyxl
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from fastapi.staticfiles import StaticFiles
from jose import JWTError, jwt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pydantic import BaseModel, Field

load_dotenv()

import client_evaluation_engine  # noqa: E402
import core.anthropic_client as anthropic_client  # noqa: E402
import gabi_engine  # noqa: E402
import repositories.client_evaluation_repository as client_evaluation_repository  # noqa: E402
import repositories.squad_repository as squad_repository  # noqa: E402
import repositories.tl_evaluation_repository as tl_evaluation_repository  # noqa: E402
import tl_evaluation_engine  # noqa: E402
import tl_evaluation_pdf  # noqa: E402
from core.oracle import is_oracle_configured, try_init_oracle_client  # noqa: E402
from db import _db  # noqa: E402

# ── Config ────────────────────────────────────────────────────────────────────

ALGORITHM = "HS256"
TOKEN_EXPIRE_HOURS = 8

USER_NOT_FOUND = "Usuário não encontrado"
SUBMISSION_NOT_FOUND = "Submission não encontrada"
MEMBER_NOT_FOUND = "Membro não encontrado"
TASK_DB_ERROR = "Erro ao consultar Task DB"
SQUAD_NOT_FOUND = "Squad não encontrada"
_SELECT_IS_ADMIN_SQL = "SELECT is_admin FROM users WHERE email=?"

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

BASE_DIR = Path(__file__).parent
DB_PATH = BASE_DIR / "users.db"
TASKS_DIR = BASE_DIR / "tasks"
_FRONTEND_DIST = BASE_DIR.parent / "frontend" / "dist"
# Fall back to the legacy monolith until frontend/dist is deployed via CI
STATIC_DIR = _FRONTEND_DIST if _FRONTEND_DIST.exists() else BASE_DIR / "static"


def _load_secret_key() -> str:
    if env := os.getenv("SECRET_KEY"):  # pragma: no cover
        return env  # pragma: no cover
    key_file = BASE_DIR / ".secret_key"
    if key_file.exists():
        return key_file.read_text().strip()
    key = secrets.token_hex(32)  # pragma: no cover
    key_file.write_text(key)  # pragma: no cover
    return key  # pragma: no cover


SECRET_KEY = _load_secret_key()

oauth2 = OAuth2PasswordBearer(tokenUrl="/auth/login")

app = FastAPI(title="G.A.B.I.", description="Gap Auditor of Billed Input", version="3.0.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


# ── Password helpers ──────────────────────────────────────────────────────────


def _hash(password: str) -> str:
    return bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()


def _verify(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode(), hashed.encode())


# ── Database ──────────────────────────────────────────────────────────────────


def _init_db() -> None:
    with _db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id              INTEGER PRIMARY KEY AUTOINCREMENT,
                email           TEXT    UNIQUE NOT NULL,
                hashed_password TEXT    NOT NULL,
                is_admin        INTEGER NOT NULL DEFAULT 0,
                is_active       INTEGER NOT NULL DEFAULT 1,
                created_at      TEXT    NOT NULL DEFAULT (datetime('now'))
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS task_uploads (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                filename    TEXT NOT NULL,
                path        TEXT NOT NULL,
                competencia TEXT NOT NULL DEFAULT '',
                uploaded_at TEXT NOT NULL DEFAULT (datetime('now')),
                uploaded_by TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS settings (
                key   TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS billing_rates (
                profile    TEXT PRIMARY KEY,
                valor_hora REAL NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS team_members (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                user_email    TEXT    UNIQUE NOT NULL,
                cpf           TEXT    NOT NULL,
                subcontrato   TEXT    NOT NULL,
                profile       TEXT    NOT NULL,
                nome_completo TEXT    NOT NULL DEFAULT '',
                is_active     INTEGER NOT NULL DEFAULT 1,
                created_at    TEXT    NOT NULL DEFAULT (datetime('now'))
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS ga_submissions (
                id                  INTEGER PRIMARY KEY AUTOINCREMENT,
                user_email          TEXT    NOT NULL,
                competencia         TEXT    NOT NULL,
                nome_profissional   TEXT    NOT NULL,
                task_id             TEXT    NOT NULL,
                horas_task          REAL    NOT NULL,
                ga_filename         TEXT    NOT NULL,
                ga_data             BLOB    NOT NULL,
                validation_snapshot TEXT    NOT NULL,
                status              TEXT    NOT NULL DEFAULT 'pending',
                rejection_reason    TEXT,
                submitted_at        TEXT    NOT NULL DEFAULT (datetime('now')),
                reviewed_at         TEXT,
                reviewed_by         TEXT,
                UNIQUE (user_email, competencia)
            )
        """)
        conn.execute(
            "INSERT OR IGNORE INTO settings (key, value) VALUES ('billing_cnpj', ?)",
            (os.getenv("BILLING_CNPJ", ""),),
        )
        conn.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('billing_email_to', '')")
        conn.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('billing_email_tag', '')")

        conn.execute("""
            CREATE TABLE IF NOT EXISTS client_evaluations (
                id                       INTEGER PRIMARY KEY AUTOINCREMENT,
                competencia              TEXT    NOT NULL,
                nome_extraido            TEXT    NOT NULL,
                user_email               TEXT,
                match_status             TEXT    NOT NULL,
                fornecedor_validado      INTEGER NOT NULL,
                comportamental_nota      INTEGER,
                comportamental_comentario TEXT,
                tecnica_nota             INTEGER,
                tecnica_comentario       TEXT,
                prazo_nota               INTEGER,
                prazo_comentario         TEXT,
                comentario_geral         TEXT,
                raw_email_text           TEXT    NOT NULL,
                email_message_id         TEXT,
                origem                   TEXT    NOT NULL,
                status                   TEXT    NOT NULL DEFAULT 'rascunho',
                created_by               TEXT    NOT NULL,
                created_at               TIMESTAMP NOT NULL DEFAULT (datetime('now')),
                reviewed_by              TEXT,
                reviewed_at              TIMESTAMP,
                edited_by                TEXT,
                edited_at                TIMESTAMP
            )
        """)
        # Partial index — prevents duplicate approved records; rascunho can coexist with aprovado
        # (required for the confirm_override re-ingest path — see TC-12).
        conn.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS idx_ce_user_competencia
              ON client_evaluations (user_email, competencia)
              WHERE user_email IS NOT NULL AND status = 'aprovado'
        """)
        conn.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS idx_ce_message_id
              ON client_evaluations (email_message_id)
              WHERE email_message_id IS NOT NULL
        """)

        # ── TL Evaluation domain tables (in dependency order) ─────────────────
        conn.execute("""
            CREATE TABLE IF NOT EXISTS squads (
                id                   INTEGER PRIMARY KEY AUTOINCREMENT,
                nome                 TEXT    NOT NULL UNIQUE,
                tech_lead_user_email TEXT    NOT NULL UNIQUE,
                created_at           TIMESTAMP NOT NULL DEFAULT (datetime('now'))
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS squad_members (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                squad_id    INTEGER NOT NULL REFERENCES squads(id),
                nome        TEXT    NOT NULL,
                papel       TEXT    NOT NULL,
                senioridade TEXT    NOT NULL,
                fornecedor  TEXT    NOT NULL,
                ativo       INTEGER NOT NULL DEFAULT 1,
                created_at  TIMESTAMP NOT NULL DEFAULT (datetime('now')),
                removed_at  TIMESTAMP
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS tl_evaluations (
                id                       INTEGER PRIMARY KEY AUTOINCREMENT,
                squad_id                 INTEGER NOT NULL REFERENCES squads(id),
                squad_member_id          INTEGER NOT NULL REFERENCES squad_members(id),
                competencia              TEXT    NOT NULL,
                comportamental_nota      INTEGER,
                comportamental_paragrafo TEXT,
                tecnica_nota             INTEGER,
                tecnica_paragrafo        TEXT,
                prazo_nota               INTEGER,
                prazo_paragrafo          TEXT,
                janela_reduzida          INTEGER NOT NULL DEFAULT 0,
                status                   TEXT    NOT NULL DEFAULT 'rascunho',
                confirmed_by             TEXT,
                created_at               TIMESTAMP NOT NULL DEFAULT (datetime('now')),
                confirmed_at             TIMESTAMP
            )
        """)
        conn.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS idx_tl_eval_member_competencia
              ON tl_evaluations (squad_id, competencia, squad_member_id)
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS squad_competencia_closures (
                squad_id    INTEGER NOT NULL REFERENCES squads(id),
                competencia TEXT    NOT NULL,
                closed_at   TIMESTAMP NOT NULL DEFAULT (datetime('now')),
                closed_by   TEXT    NOT NULL
            )
        """)
        conn.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS idx_sqc_closure
              ON squad_competencia_closures (squad_id, competencia)
        """)

        # Migrations for existing installs
        for ddl in [
            "ALTER TABLE task_uploads ADD COLUMN competencia TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE users ADD COLUMN is_admin INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE ga_submissions ADD COLUMN approved_snapshot TEXT",
            "ALTER TABLE team_members ADD COLUMN deactivated_from_competencia TEXT",
            "ALTER TABLE team_members ADD COLUMN nome_completo TEXT NOT NULL DEFAULT ''",
        ]:
            with contextlib.suppress(sqlite3.OperationalError):
                conn.execute(ddl)

        if not conn.execute("SELECT 1 FROM users LIMIT 1").fetchone():
            default_email = os.getenv("ADMIN_EMAIL", "tiagohlatki@db1.com.br")
            default_pass = os.getenv("ADMIN_PASSWORD", "gabi2025")
            conn.execute(
                "INSERT INTO users (email, hashed_password, is_admin) VALUES (?, ?, 1)",
                (default_email, _hash(default_pass)),
            )
        else:  # pragma: no cover — existing-install migration path
            admin_email = os.getenv("ADMIN_EMAIL", "tiagohlatki@db1.com.br")
            conn.execute(
                "UPDATE users SET is_admin=1 WHERE email=? AND is_admin=0",
                (admin_email,),
            )


# ── Auth helpers ──────────────────────────────────────────────────────────────


def _create_token(email: str) -> str:
    exp = datetime.now(tz=UTC) + timedelta(hours=TOKEN_EXPIRE_HOURS)
    return jwt.encode({"sub": email, "exp": exp}, SECRET_KEY, algorithm=ALGORITHM)


def _current_user(token: Annotated[str, Depends(oauth2)]) -> str:
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        email: str = payload.get("sub", "")
        if not email:
            raise ValueError
    except (JWTError, ValueError):  # fmt: skip
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Token inválido ou expirado",
        ) from None
    with _db() as conn:
        row = conn.execute("SELECT 1 FROM users WHERE email=? AND is_active=1", (email,)).fetchone()
    if not row:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Usuário inativo")
    return email


def _admin_user(email: Annotated[str, Depends(_current_user)]) -> str:
    with _db() as conn:
        row = conn.execute(_SELECT_IS_ADMIN_SQL, (email,)).fetchone()
    if not row or not row["is_admin"]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Acesso restrito a administradores",
        )
    return email


CurrentUser = Annotated[str, Depends(_current_user)]
AdminUser = Annotated[str, Depends(_admin_user)]


# ── Routes: auth ──────────────────────────────────────────────────────────────


@app.post("/auth/login")
async def login(form: Annotated[OAuth2PasswordRequestForm, Depends()]):
    with _db() as conn:
        row = conn.execute(
            "SELECT hashed_password FROM users WHERE email=? AND is_active=1",
            (form.username,),
        ).fetchone()
    if not row or not _verify(form.password, row["hashed_password"]):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED, detail="Credenciais inválidas"
        )
    return {"access_token": _create_token(form.username), "token_type": "bearer"}


@app.get("/auth/me")
async def me(email: CurrentUser):
    with _db() as conn:
        row = conn.execute(_SELECT_IS_ADMIN_SQL, (email,)).fetchone()
    return {"email": email, "is_admin": bool(row["is_admin"]) if row else False}


@app.post(
    "/auth/change-password",
    responses={400: {"description": "Senha atual incorreta ou nova senha inválida"}},
)
async def change_password(current_password: str, new_password: str, email: CurrentUser):
    with _db() as conn:
        row = conn.execute("SELECT hashed_password FROM users WHERE email=?", (email,)).fetchone()
        if not _verify(current_password, row["hashed_password"]):
            raise HTTPException(status_code=400, detail="Senha atual incorreta")
        if len(new_password) < 6:
            raise HTTPException(
                status_code=400, detail="Nova senha deve ter pelo menos 6 caracteres"
            )
        conn.execute(
            "UPDATE users SET hashed_password=? WHERE email=?",
            (_hash(new_password), email),
        )
    return {"ok": True}


# ── Routes: admin — users ─────────────────────────────────────────────────────


class CreateUserRequest(BaseModel):
    email: str
    password: str
    is_admin: bool = False


@app.post(
    "/api/admin/users",
    status_code=201,
    responses={
        400: {"description": "Senha inválida"},
        409: {"description": "E-mail já cadastrado"},
    },
)
async def admin_create_user(body: CreateUserRequest, _: AdminUser):
    if len(body.password) < 6:
        raise HTTPException(status_code=400, detail="Senha deve ter pelo menos 6 caracteres")
    try:
        with _db() as conn:
            conn.execute(
                "INSERT INTO users (email, hashed_password, is_admin) VALUES (?, ?, ?)",
                (body.email, _hash(body.password), int(body.is_admin)),
            )
    except sqlite3.IntegrityError as exc:
        raise HTTPException(status_code=409, detail="E-mail já cadastrado") from exc
    return {"ok": True, "email": body.email}


@app.get("/api/admin/users")
async def admin_list_users(_: AdminUser):
    with _db() as conn:
        rows = conn.execute(
            "SELECT id, email, is_admin, is_active, created_at FROM users ORDER BY id"
        ).fetchall()
    return {"users": [dict(r) for r in rows]}


@app.patch(
    "/api/admin/users/{user_id}/toggle",
    responses={
        404: {"description": "Usuário não encontrado"},
        400: {"description": "Não é possível alterar o próprio usuário"},
    },
)
async def admin_toggle_user(user_id: int, admin_email: AdminUser):
    with _db() as conn:
        row = conn.execute("SELECT email, is_active FROM users WHERE id=?", (user_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail=USER_NOT_FOUND)
        if row["email"] == admin_email:
            raise HTTPException(status_code=400, detail="Não é possível alterar o próprio usuário")
        new_active = 0 if row["is_active"] else 1
        conn.execute("UPDATE users SET is_active=? WHERE id=?", (new_active, user_id))
    return {"ok": True, "is_active": bool(new_active)}


@app.delete(
    "/api/admin/users/{user_id}",
    responses={
        404: {"description": "Usuário não encontrado"},
        400: {"description": "Não é possível remover o próprio usuário"},
    },
)
async def admin_delete_user(user_id: int, admin_email: AdminUser):
    with _db() as conn:
        row = conn.execute("SELECT email FROM users WHERE id=?", (user_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail=USER_NOT_FOUND)
        if row["email"] == admin_email:
            raise HTTPException(status_code=400, detail="Não é possível remover o próprio usuário")
        conn.execute("DELETE FROM users WHERE id=?", (user_id,))
    return {"ok": True}


# ── Routes: task status (public) ──────────────────────────────────────────────


@app.get("/api/task/status")
async def task_status():
    with _db() as conn:
        row = conn.execute(
            "SELECT id, filename, competencia, uploaded_at, uploaded_by FROM task_uploads "
            "ORDER BY competencia DESC, id DESC LIMIT 1"
        ).fetchone()
    if not row:
        return {"available": False}
    return {"available": True, **dict(row)}


@app.get("/api/task/competencias")
async def task_competencias():
    with _db() as conn:
        rows = conn.execute("""
            SELECT t.competencia, t.uploaded_at, t.uploaded_by
            FROM task_uploads t
            WHERE t.id IN (SELECT MAX(id) FROM task_uploads GROUP BY competencia)
            ORDER BY t.competencia DESC
        """).fetchall()
    return {"competencias": [dict(r) for r in rows]}


# ── Routes: task entries report ────────────────────────────────────────────────


@app.get("/api/task/default-project")
async def task_default_project():
    from repositories.task_repository import get_default_project

    return {"project": get_default_project()}


_TASK_REPORT_HEADERS = ["Projeto", "Tarefa", "Descrição", "Início", "Término", "Horas", "Área"]
_TASK_REPORT_COLUMN_WIDTHS = [14, 10, 42, 10, 10, 10, 16]
_INDENT_COLUMNS = 1  # blank leftmost column, used only as a visual margin
_INDENT_COLUMN_WIDTH = 3
_DESCRIPTION_COLUMN = 3 + _INDENT_COLUMNS

_DAY_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
_COLUMN_HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
_UNEXPECTED_AREA_FILL = PatternFill(start_color="FDE68A", end_color="FDE68A", fill_type="solid")
_THIN_SIDE = Side(style="thin", color="CBD5E1")
_THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
# Projeto, Tarefa, Início, Término, Horas, Área
_CENTERED_ENTRY_COLUMNS = tuple(c + _INDENT_COLUMNS for c in (1, 2, 4, 5, 6, 7))
_REPORT_TIMEZONE = ZoneInfo("America/Sao_Paulo")


def _append_indented_row(ws, values: list) -> None:
    ws.append([""] * _INDENT_COLUMNS + values)


def _fetch_task_report(responsavel: str, sd: date, ed: date) -> list[dict]:
    from repositories.task_repository import get_task_entries

    entries = get_task_entries(responsavel, sd, ed)
    return gabi_engine.build_task_report(entries)


async def _task_report_days(email: str, start_date: str, end_date: str) -> list[dict]:
    from repositories.task_repository import responsavel_from_email

    sd, ed = date.fromisoformat(start_date), date.fromisoformat(end_date)
    responsavel = responsavel_from_email(email)
    fetch = functools.partial(_fetch_task_report, responsavel, sd, ed)
    try:
        return await anyio.to_thread.run_sync(fetch)
    except Exception as exc:
        raise HTTPException(status_code=503, detail=f"{TASK_DB_ERROR}: {exc}") from exc


def _apply_task_report_filters(
    days: list[dict], only_faturado: bool, only_default_project: bool
) -> list[dict]:
    from repositories.task_repository import get_default_project

    project = get_default_project() if only_default_project else None
    return gabi_engine.filter_task_report(days, only_faturado, project)


def _append_bold_labeled_row(ws, label: str, value: str) -> None:
    _append_indented_row(ws, [label, value])
    ws.cell(row=ws.max_row, column=1 + _INDENT_COLUMNS).font = Font(bold=True)


def _write_task_report_summary(
    ws, email: str, start_date: str, end_date: str, total: float
) -> None:
    sd, ed = date.fromisoformat(start_date), date.fromisoformat(end_date)
    periodo = f"{sd:%d/%m/%Y} a {ed:%d/%m/%Y}"
    _append_indented_row(ws, ["Relatório de Apontamentos Task"])
    ws.cell(row=ws.max_row, column=1 + _INDENT_COLUMNS).font = Font(bold=True, size=14)
    ws.append([])
    _append_bold_labeled_row(ws, "Usuário:", email.upper())
    _append_bold_labeled_row(ws, "Período:", periodo)
    _append_bold_labeled_row(ws, "Total do período:", f"{total:.2f}h")
    ws.append([])


def _write_task_report_day_header(ws, day: dict) -> None:
    _append_indented_row(
        ws, [day["date"], "", "", "", "", f"Total do dia: {day['total']:.2f}h", ""]
    )
    row = ws.max_row
    start = 1 + _INDENT_COLUMNS
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=start + 4)
    ws.merge_cells(start_row=row, start_column=start + 5, end_row=row, end_column=start + 6)
    for cell in ws[row][_INDENT_COLUMNS:]:
        cell.fill = _DAY_HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
    ws.cell(row=row, column=start + 5).alignment = Alignment(horizontal="right")


def _write_task_report_column_header(ws) -> None:
    _append_indented_row(ws, list(_TASK_REPORT_HEADERS))
    row = ws.max_row
    for cell in ws[row][_INDENT_COLUMNS:]:
        cell.fill = _COLUMN_HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = _THIN_BORDER
        cell.alignment = Alignment(horizontal="center")


def _write_task_report_entry(ws, entry: dict) -> None:
    _append_indented_row(
        ws,
        [
            entry["project"],
            entry["task_code"],
            entry["description"],
            entry["start"],
            entry["end"],
            entry["hours"],
            entry["area"],
        ],
    )
    row = ws.max_row
    for cell in ws[row][_INDENT_COLUMNS:]:
        cell.border = _THIN_BORDER
    for col in _CENTERED_ENTRY_COLUMNS:
        ws.cell(row=row, column=col).alignment = Alignment(horizontal="center")
    ws.cell(row=row, column=6 + _INDENT_COLUMNS).number_format = "0.00"
    if entry["unexpected_area"]:
        ws.cell(row=row, column=7 + _INDENT_COLUMNS).fill = _UNEXPECTED_AREA_FILL


def _write_task_report_generation_footer(ws, generated_by: str) -> None:
    generated_at = datetime.now(tz=_REPORT_TIMEZONE).strftime("%d/%m/%Y %H:%M:%S")
    ws.append([])
    _append_indented_row(ws, [f"Gerado por {generated_by} em {generated_at}"])


def _description_column_width(days: list[dict]) -> int:
    header_len = len(_TASK_REPORT_HEADERS[2])  # "Descrição"
    max_len = max(
        (len(entry["description"]) for day in days for entry in day["entries"]), default=0
    )
    return max(header_len, max_len) + 2


def _fill_task_report_worksheet(
    ws, email: str, start_date: str, end_date: str, days: list[dict], generated_by: str
) -> None:
    ws.sheet_view.showGridLines = False
    total = round(sum(day["total"] for day in days), 2)
    _write_task_report_summary(ws, email, start_date, end_date, total)

    for day in days:
        _write_task_report_day_header(ws, day)
        _write_task_report_column_header(ws)
        for entry in day["entries"]:
            _write_task_report_entry(ws, entry)
        ws.append([])

    _write_task_report_generation_footer(ws, generated_by)

    ws.column_dimensions[get_column_letter(_INDENT_COLUMNS)].width = _INDENT_COLUMN_WIDTH
    for i, width in enumerate(_TASK_REPORT_COLUMN_WIDTHS, start=1 + _INDENT_COLUMNS):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.column_dimensions[get_column_letter(_DESCRIPTION_COLUMN)].width = _description_column_width(
        days
    )


def _task_report_xlsx_response(
    days: list[dict], email: str, start_date: str, end_date: str, filename: str, generated_by: str
) -> Response:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Apontamentos"
    _fill_task_report_worksheet(ws, email, start_date, end_date, days, generated_by)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get(
    "/api/task/my-entries",
    responses={503: {"description": TASK_DB_ERROR}},
)
async def my_task_entries(start_date: str, end_date: str, email: CurrentUser = ""):
    days = await _task_report_days(email, start_date, end_date)
    return {"days": days}


@app.get(
    "/api/task/my-entries/export",
    responses={503: {"description": TASK_DB_ERROR}},
)
async def export_my_task_entries(
    start_date: str,
    end_date: str,
    only_faturado: bool = True,
    only_default_project: bool = True,
    email: CurrentUser = "",
):
    days = await _task_report_days(email, start_date, end_date)
    days = _apply_task_report_filters(days, only_faturado, only_default_project)
    filename = f"apontamentos_{start_date}_{end_date}.xlsx"
    return _task_report_xlsx_response(days, email, start_date, end_date, filename, email)


# ── Routes: oracle status (public) ────────────────────────────────────────────


@app.get("/api/oracle/status")
async def oracle_status():
    if not is_oracle_configured():
        return {"available": False}
    from repositories.task_repository import ping  # pragma: no cover

    available = False  # pragma: no cover
    with anyio.move_on_after(3), contextlib.suppress(Exception):  # pragma: no cover
        available = await anyio.to_thread.run_sync(ping, cancellable=True)  # pragma: no cover
    return {"available": available}  # pragma: no cover


# ── Routes: admin task management ─────────────────────────────────────────────


@app.post("/api/admin/task")
async def admin_upload_task(
    task: Annotated[UploadFile, File()],
    competencia: Annotated[str, Form()],
    email: AdminUser,
):
    TASKS_DIR.mkdir(exist_ok=True)
    ts = datetime.now(tz=UTC).strftime("%Y%m%d_%H%M%S")
    safe_name = f"{ts}_{task.filename or 'task.xlsx'}"
    dest = TASKS_DIR / safe_name
    async with await anyio.open_file(dest, "wb") as f:
        await f.write(await task.read())
    with _db() as conn:
        conn.execute(
            "INSERT INTO task_uploads (filename, path, competencia, uploaded_by)"
            " VALUES (?, ?, ?, ?)",
            (task.filename or "task.xlsx", str(dest), competencia, email),
        )
    return {"ok": True, "filename": task.filename, "competencia": competencia}


@app.delete(
    "/api/admin/task/{task_id}",
    responses={404: {"description": "Upload não encontrado"}},
)
async def admin_delete_task(task_id: int, _: AdminUser):
    with _db() as conn:
        row = conn.execute("SELECT path FROM task_uploads WHERE id=?", (task_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Upload não encontrado")
        conn.execute("DELETE FROM task_uploads WHERE id=?", (task_id,))
    with contextlib.suppress(Exception):
        Path(row["path"]).unlink(missing_ok=True)
    return {"ok": True}


@app.get("/api/admin/tasks")
async def admin_list_tasks(_: AdminUser):
    with _db() as conn:
        rows = conn.execute(
            "SELECT id, filename, competencia, uploaded_at, uploaded_by"
            " FROM task_uploads ORDER BY id DESC"
        ).fetchall()
    return {"tasks": [dict(r) for r in rows]}


NOT_TEAM_MEMBER = "Usuário não faz parte do time"


def _require_team_member(email: str) -> None:
    with _db() as conn:
        row = conn.execute("SELECT 1 FROM team_members WHERE user_email=?", (email,)).fetchone()
    if not row:
        raise HTTPException(status_code=403, detail=NOT_TEAM_MEMBER)


@app.get(
    "/api/admin/task/entries",
    responses={403: {"description": NOT_TEAM_MEMBER}, 503: {"description": TASK_DB_ERROR}},
)
async def admin_task_entries(email: str, start_date: str, end_date: str, _: AdminUser):
    _require_team_member(email)
    days = await _task_report_days(email, start_date, end_date)
    return {"days": days}


@app.get(
    "/api/admin/task/entries/export",
    responses={403: {"description": NOT_TEAM_MEMBER}, 503: {"description": TASK_DB_ERROR}},
)
async def admin_export_task_entries(
    email: str,
    start_date: str,
    end_date: str,
    admin_email: AdminUser,
    only_faturado: bool = True,
    only_default_project: bool = True,
):
    _require_team_member(email)
    days = await _task_report_days(email, start_date, end_date)
    days = _apply_task_report_filters(days, only_faturado, only_default_project)
    filename = f"apontamentos_{email}_{start_date}_{end_date}.xlsx"
    return _task_report_xlsx_response(days, email, start_date, end_date, filename, admin_email)


# ── Routes: validate ──────────────────────────────────────────────────────────


def _resolve_task_file(competencia: str | None) -> str:
    with _db() as conn:
        if competencia:
            row = conn.execute(
                "SELECT path FROM task_uploads WHERE competencia=? ORDER BY id DESC LIMIT 1",
                (competencia,),
            ).fetchone()
        else:  # pragma: no cover — fallback used only by Oracle rerun, not in file-mode tests
            row = conn.execute(
                "SELECT path FROM task_uploads ORDER BY competencia DESC, id DESC LIMIT 1"
            ).fetchone()  # pragma: no cover
    if not row:
        raise ValueError(
            "Nenhum relatório Task disponível. Aguarde o administrador fazer o upload."
        )
    if not Path(row["path"]).exists():
        raise ValueError(
            "Arquivo Task não encontrado no servidor. Entre em contato com o administrador."
        )
    return row["path"]


def _process_ga_with_oracle(
    ga_path: str, email: str, sd: date, ed: date, year: int, month: int
) -> tuple[dict | None, str | None]:
    """Returns (person, error_message). One of them is always None."""
    from repositories.task_repository import get_task_hours, responsavel_from_email

    try:
        ga_name, ga_daily, _, _ = gabi_engine.parse_ga(ga_path)
    except Exception as exc:
        return None, f"Erro ao ler GA: {exc}"

    try:
        responsavel = responsavel_from_email(email)
        task_daily, task_nf_daily = get_task_hours(responsavel, sd, ed)
    except Exception as exc:
        return None, f"Erro ao consultar Task DB: {exc}"

    person = gabi_engine.process_with_task_data(
        responsavel, task_daily, ga_name, ga_daily, year, month, sd, ed, task_nf_daily
    )
    return person, None


async def _save_ga_upload(ga_file: UploadFile, tmpdir: str) -> str:
    ga_path = os.path.join(tmpdir, ga_file.filename or "ga.xls")
    async with await anyio.open_file(ga_path, "wb") as f:
        await f.write(await ga_file.read())
    return ga_path


async def _run_oracle(
    ga_path: str, email: str, sd: date, ed: date, year: int, month: int
) -> tuple[dict | None, str | None]:
    try:
        return await anyio.to_thread.run_sync(
            functools.partial(_process_ga_with_oracle, ga_path, email, sd, ed, year, month)
        )
    except Exception as exc:
        return None, f"Erro interno ao processar: {exc}"


def _run_file(ga_path: str, competencia: str | None) -> tuple[list, list]:
    task_path = _resolve_task_file(competencia)
    return gabi_engine.process(task_path, [ga_path])


@app.post(
    "/api/validate",
    responses={422: {"description": "Erro de validação ou arquivo Task ausente"}},
)
async def validate(
    ga: Annotated[list[UploadFile], File()],
    start_date: Annotated[str, Form()],
    end_date: Annotated[str, Form()],
    use_file_fallback: Annotated[bool, Form()] = False,
    competencia: Annotated[str | None, Form()] = None,
    email: CurrentUser = "",
):
    sd = date.fromisoformat(start_date)
    ed = date.fromisoformat(end_date)
    year, month = sd.year, sd.month
    use_oracle = is_oracle_configured() and not use_file_fallback
    persons: list = []
    errors: list = []

    with tempfile.TemporaryDirectory() as tmpdir:
        for ga_file in ga:
            ga_path = await _save_ga_upload(ga_file, tmpdir)
            if use_oracle:
                person, error = await _run_oracle(ga_path, email, sd, ed, year, month)
                if error:
                    errors.append(error)
                else:
                    persons.append(person)
            else:
                try:
                    p, e = _run_file(ga_path, competencia)
                    persons.extend(p)
                    errors.extend(e)
                except Exception as exc:
                    raise HTTPException(status_code=422, detail=str(exc)) from exc

    if not persons and errors:
        raise HTTPException(status_code=422, detail="\n".join(errors))

    return {"persons": persons, "errors": errors}


# ── Helpers ───────────────────────────────────────────────────────────────────

_MONTH_ABBR_PT = [
    "jan",
    "fev",
    "mar",
    "abr",
    "mai",
    "jun",
    "jul",
    "ago",
    "set",
    "out",
    "nov",
    "dez",
]


def _competencia_label(competencia: str) -> str:
    year, month = competencia.split("-")
    return f"{_MONTH_ABBR_PT[int(month) - 1]}/{year[2:]}"


def _competencia_date_range(competencia: str) -> tuple[date, date]:
    y, m = (int(x) for x in competencia.split("-"))
    last_day = (date(y, m % 12 + 1, 1) - timedelta(days=1)).day if m < 12 else 31
    return date(y, m, 1), date(y, m, last_day)


# ── Routes: settings (admin) ──────────────────────────────────────────────────


@app.get("/api/admin/settings")
async def admin_get_settings(_: AdminUser):
    with _db() as conn:
        rows = conn.execute("SELECT key, value FROM settings").fetchall()
    return {r["key"]: r["value"] for r in rows}


class UpdateSettingRequest(BaseModel):
    value: str


@app.put("/api/admin/settings/{key}")
async def admin_update_setting(key: str, body: UpdateSettingRequest, _: AdminUser):
    with _db() as conn:
        conn.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?)"
            " ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (key, body.value),
        )
    return {"ok": True}


# ── Routes: billing rates (admin) ─────────────────────────────────────────────


class BillingRateRequest(BaseModel):
    profile: str
    valor_hora: float


@app.get("/api/admin/billing-rates")
async def admin_list_billing_rates(_: AdminUser):
    with _db() as conn:
        rows = conn.execute(
            "SELECT profile, valor_hora FROM billing_rates ORDER BY profile"
        ).fetchall()
    return {"rates": [dict(r) for r in rows]}


@app.post(
    "/api/admin/billing-rates",
    status_code=201,
    responses={409: {"description": "Perfil já cadastrado"}},
)
async def admin_create_billing_rate(body: BillingRateRequest, _: AdminUser):
    try:
        with _db() as conn:
            conn.execute(
                "INSERT INTO billing_rates (profile, valor_hora) VALUES (?, ?)",
                (body.profile, body.valor_hora),
            )
    except sqlite3.IntegrityError as exc:
        raise HTTPException(status_code=409, detail="Perfil já cadastrado") from exc
    return {"ok": True, "profile": body.profile}


class UpdateBillingRateRequest(BaseModel):
    valor_hora: float


@app.put(
    "/api/admin/billing-rates/{profile}",
    responses={404: {"description": "Perfil não encontrado"}},
)
async def admin_update_billing_rate(profile: str, body: UpdateBillingRateRequest, _: AdminUser):
    with _db() as conn:
        result = conn.execute(
            "UPDATE billing_rates SET valor_hora=? WHERE profile=?",
            (body.valor_hora, profile),
        )
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="Perfil não encontrado")
    return {"ok": True}


@app.delete(
    "/api/admin/billing-rates/{profile}",
    responses={404: {"description": "Perfil não encontrado"}},
)
async def admin_delete_billing_rate(profile: str, _: AdminUser):
    with _db() as conn:
        result = conn.execute("DELETE FROM billing_rates WHERE profile=?", (profile,))
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="Perfil não encontrado")
    return {"ok": True}


# ── Routes: team members (admin) ──────────────────────────────────────────────


class CreateTeamMemberRequest(BaseModel):
    user_email: str
    cpf: str
    subcontrato: str
    profile: str
    nome_completo: str = Field(min_length=1)


class UpdateTeamMemberRequest(BaseModel):
    cpf: str
    subcontrato: str
    profile: str
    nome_completo: str = Field(min_length=1)
    is_active: bool = True


@app.get("/api/admin/team")
async def admin_list_team(_: AdminUser):
    with _db() as conn:
        rows = conn.execute("""
            SELECT id, user_email, cpf, subcontrato, profile, nome_completo, is_active, created_at
            FROM team_members ORDER BY id
        """).fetchall()
    return {"members": [dict(r) for r in rows]}


@app.post(
    "/api/admin/team",
    status_code=201,
    responses={
        404: {"description": USER_NOT_FOUND},
        400: {"description": "Perfil não cadastrado em tarifas"},
        409: {"description": "Usuário já é membro do time"},
    },
)
async def admin_create_team_member(body: CreateTeamMemberRequest, _: AdminUser):
    with _db() as conn:
        user = conn.execute("SELECT 1 FROM users WHERE email=?", (body.user_email,)).fetchone()
        if not user:
            raise HTTPException(status_code=404, detail=USER_NOT_FOUND)
        rate = conn.execute(
            "SELECT 1 FROM billing_rates WHERE profile=?", (body.profile,)
        ).fetchone()
        if not rate:
            raise HTTPException(status_code=400, detail="Perfil não cadastrado em tarifas")
        try:
            conn.execute(
                "INSERT INTO team_members (user_email, cpf, subcontrato, profile, nome_completo)"
                " VALUES (?, ?, ?, ?, ?)",
                (body.user_email, body.cpf, body.subcontrato, body.profile, body.nome_completo),
            )
        except sqlite3.IntegrityError as exc:
            raise HTTPException(status_code=409, detail="Usuário já é membro do time") from exc
    return {"ok": True}


@app.put(
    "/api/admin/team/{member_id}",
    responses={
        400: {"description": "Perfil não cadastrado em tarifas"},
        404: {"description": MEMBER_NOT_FOUND},
    },
)
async def admin_update_team_member(member_id: int, body: UpdateTeamMemberRequest, _: AdminUser):
    with _db() as conn:
        rate = conn.execute(
            "SELECT 1 FROM billing_rates WHERE profile=?", (body.profile,)
        ).fetchone()
        if not rate:
            raise HTTPException(status_code=400, detail="Perfil não cadastrado em tarifas")
        result = conn.execute(
            "UPDATE team_members SET cpf=?, subcontrato=?, profile=?, nome_completo=?, is_active=?"
            " WHERE id=?",
            (
                body.cpf,
                body.subcontrato,
                body.profile,
                body.nome_completo,
                int(body.is_active),
                member_id,
            ),
        )
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail=MEMBER_NOT_FOUND)
    return {"ok": True}


@app.delete(
    "/api/admin/team/{member_id}",
    responses={404: {"description": MEMBER_NOT_FOUND}},
)
async def admin_delete_team_member(member_id: int, _: AdminUser):
    with _db() as conn:
        result = conn.execute("DELETE FROM team_members WHERE id=?", (member_id,))
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail=MEMBER_NOT_FOUND)
    return {"ok": True}


class DeactivateMemberRequest(BaseModel):
    deactivated_from_competencia: str


@app.patch(
    "/api/admin/team/{member_id}/deactivate",
    responses={404: {"description": MEMBER_NOT_FOUND}},
)
async def admin_deactivate_team_member(member_id: int, body: DeactivateMemberRequest, _: AdminUser):
    with _db() as conn:
        result = conn.execute(
            "UPDATE team_members SET is_active=0, deactivated_from_competencia=? WHERE id=?",
            (body.deactivated_from_competencia, member_id),
        )
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail=MEMBER_NOT_FOUND)
    return {"ok": True}


@app.patch(
    "/api/admin/team/{member_id}/activate",
    responses={404: {"description": MEMBER_NOT_FOUND}},
)
async def admin_activate_team_member(member_id: int, _: AdminUser):
    with _db() as conn:
        result = conn.execute(
            "UPDATE team_members SET is_active=1, deactivated_from_competencia=NULL WHERE id=?",
            (member_id,),
        )
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail=MEMBER_NOT_FOUND)
    return {"ok": True}


# ── Routes: submissions (user) ────────────────────────────────────────────────


@app.post("/api/submissions", status_code=201)
async def submit_validation(
    ga: Annotated[UploadFile, File()],
    competencia: Annotated[str, Form()],
    nome_profissional: Annotated[str, Form()],
    task_id: Annotated[str, Form()],
    horas_task: Annotated[float, Form()],
    snapshot: Annotated[str, Form()],
    email: CurrentUser = "",
):
    ga_bytes = await ga.read()
    with _db() as conn:
        conn.execute(
            """
            INSERT INTO ga_submissions
                (user_email, competencia, nome_profissional, task_id, horas_task,
                 ga_filename, ga_data, validation_snapshot, status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'pending')
            ON CONFLICT(user_email, competencia) DO UPDATE SET
                nome_profissional = excluded.nome_profissional,
                task_id           = excluded.task_id,
                horas_task        = excluded.horas_task,
                ga_filename       = excluded.ga_filename,
                ga_data           = excluded.ga_data,
                validation_snapshot = excluded.validation_snapshot,
                status            = 'pending',
                rejection_reason  = NULL,
                submitted_at      = datetime('now'),
                reviewed_at       = NULL,
                reviewed_by       = NULL
            """,
            (
                email,
                competencia,
                nome_profissional,
                task_id,
                horas_task,
                ga.filename or "ga.xls",
                ga_bytes,
                snapshot,
            ),
        )
    return {"ok": True}


@app.get("/api/submissions/status")
async def submission_status(competencia: str, email: CurrentUser = ""):
    with _db() as conn:
        row = conn.execute(
            "SELECT id, status, rejection_reason, submitted_at FROM ga_submissions "
            "WHERE user_email=? AND competencia=?",
            (email, competencia),
        ).fetchone()
    if not row:
        return {"submitted": False}
    return {"submitted": True, **dict(row)}


@app.get("/api/submissions/my")
async def list_my_submissions(email: CurrentUser = ""):
    with _db() as conn:
        rows = conn.execute(
            """
            SELECT id, competencia, nome_profissional, task_id, horas_task,
                   ga_filename, status, rejection_reason, submitted_at, reviewed_at
            FROM ga_submissions
            WHERE user_email=?
            ORDER BY competencia DESC
            """,
            (email,),
        ).fetchall()
    return {"submissions": [dict(r) for r in rows]}


@app.get(
    "/api/submissions/{submission_id}",
    responses={404: {"description": "Submissão não encontrada"}},
)
async def get_my_submission(submission_id: int, email: CurrentUser = ""):
    with _db() as conn:
        row = conn.execute(
            """
            SELECT id, competencia, nome_profissional, task_id, horas_task,
                   ga_filename, status, rejection_reason, submitted_at, reviewed_at,
                   validation_snapshot, approved_snapshot
            FROM ga_submissions
            WHERE id=? AND user_email=?
            """,
            (submission_id, email),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Submissão não encontrada")

    data = dict(row)
    status_val = data["status"]

    if status_val == "approved":
        raw = data.get("approved_snapshot")
        if raw:
            snapshot = json.loads(raw)
            snapshot_source = "approved"
        else:  # pragma: no cover — approved submission with null snapshot (legacy edge case)
            snapshot = None  # pragma: no cover
            snapshot_source = None  # pragma: no cover
    else:
        raw = data.get("validation_snapshot")
        snapshot = json.loads(raw) if raw else None
        snapshot_source = "original" if snapshot else None

    return {
        "id": data["id"],
        "competencia": data["competencia"],
        "nome_profissional": data["nome_profissional"],
        "task_id": data["task_id"],
        "horas_task": data["horas_task"],
        "ga_filename": data["ga_filename"],
        "status": status_val,
        "rejection_reason": data["rejection_reason"],
        "submitted_at": data["submitted_at"],
        "reviewed_at": data["reviewed_at"],
        "snapshot": snapshot,
        "snapshot_source": snapshot_source,
    }


# ── Routes: submissions (admin) ───────────────────────────────────────────────


@app.get("/api/admin/submissions/summary")
async def admin_submissions_summary(_: AdminUser):
    with _db() as conn:
        members = conn.execute("SELECT deactivated_from_competencia FROM team_members").fetchall()
        rows = conn.execute("""
            SELECT
                competencia,
                COUNT(*) AS total_submitted,
                SUM(CASE WHEN status='approved' THEN 1 ELSE 0 END) AS total_approved,
                SUM(CASE WHEN status='pending'  THEN 1 ELSE 0 END) AS total_pending,
                SUM(CASE WHEN status='rejected' THEN 1 ELSE 0 END) AS total_rejected
            FROM ga_submissions
            GROUP BY competencia
            ORDER BY competencia DESC
        """).fetchall()

    def members_for_competencia(comp: str) -> int:
        return sum(
            1
            for m in members
            if m["deactivated_from_competencia"] is None or m["deactivated_from_competencia"] > comp
        )

    total_active = sum(1 for m in members if m["deactivated_from_competencia"] is None)
    competencias = []
    for row in rows:
        entry = dict(row)
        entry["total_members"] = members_for_competencia(row["competencia"])
        competencias.append(entry)
    return {"total_members": total_active, "competencias": competencias}


@app.get("/api/admin/submissions")
async def admin_list_submissions(competencia: str, _: AdminUser):
    with _db() as conn:
        rows = conn.execute(
            """
            SELECT s.id, s.user_email, s.competencia, s.nome_profissional, s.task_id, s.horas_task,
                   s.ga_filename, s.status, s.rejection_reason, s.submitted_at, s.reviewed_at,
                   s.reviewed_by, COALESCE(tm.subcontrato, '') AS subcontrato
            FROM ga_submissions s
            LEFT JOIN team_members tm ON tm.user_email = s.user_email
            WHERE s.competencia=?
            ORDER BY s.nome_profissional
        """,
            (competencia,),
        ).fetchall()
    return {"submissions": [dict(r) for r in rows]}


@app.get("/api/admin/submissions/{competencia}/missing")
async def admin_missing_submissions(competencia: str, _: AdminUser):
    with _db() as conn:
        members = conn.execute(
            "SELECT user_email, deactivated_from_competencia FROM team_members"
        ).fetchall()
        submitted_emails = {
            row["user_email"]
            for row in conn.execute(
                "SELECT user_email FROM ga_submissions WHERE competencia=?",
                (competencia,),
            ).fetchall()
        }
    missing = [
        {"user_email": m["user_email"]}
        for m in members
        if (
            m["deactivated_from_competencia"] is None
            or m["deactivated_from_competencia"] > competencia
        )
        and m["user_email"] not in submitted_emails
    ]
    return {"missing": missing}


class ApproveSubmissionRequest(BaseModel):
    horas_task: float
    approved_snapshot: str


@app.patch(
    "/api/admin/submissions/{submission_id}/approve",
    responses={404: {"description": SUBMISSION_NOT_FOUND}},
)
async def admin_approve_submission(
    submission_id: int, admin_email: AdminUser, body: ApproveSubmissionRequest | None = None
):
    with _db() as conn:
        if body is None:
            row = conn.execute(
                "SELECT horas_task, validation_snapshot FROM ga_submissions WHERE id=?",
                (submission_id,),
            ).fetchone()
            if not row:
                raise HTTPException(status_code=404, detail=SUBMISSION_NOT_FOUND)
            horas_task = row["horas_task"]
            approved_snapshot = row["validation_snapshot"]
        else:
            horas_task = body.horas_task
            approved_snapshot = body.approved_snapshot
        result = conn.execute(
            "UPDATE ga_submissions SET status='approved', rejection_reason=NULL, "
            "horas_task=?, approved_snapshot=?, "
            "reviewed_at=datetime('now'), reviewed_by=? WHERE id=?",
            (horas_task, approved_snapshot, admin_email, submission_id),
        )
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail=SUBMISSION_NOT_FOUND)
    return {"ok": True}


class RejectSubmissionRequest(BaseModel):
    reason: str = ""


@app.patch(
    "/api/admin/submissions/{submission_id}/reject",
    responses={404: {"description": SUBMISSION_NOT_FOUND}},
)
async def admin_reject_submission(
    submission_id: int, body: RejectSubmissionRequest, admin_email: AdminUser
):
    with _db() as conn:
        result = conn.execute(
            "UPDATE ga_submissions SET status='rejected', rejection_reason=?, "
            "reviewed_at=datetime('now'), reviewed_by=? WHERE id=?",
            (body.reason, admin_email, submission_id),
        )
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail=SUBMISSION_NOT_FOUND)
    return {"ok": True}


@app.post(
    "/api/admin/submissions/{submission_id}/rerun",
    responses={
        404: {"description": SUBMISSION_NOT_FOUND},
        422: {"description": "Erro ao reprocessar"},
    },
)
async def admin_rerun_submission(submission_id: int, _: AdminUser):
    with _db() as conn:
        row = conn.execute(
            "SELECT user_email, competencia, ga_filename, ga_data FROM ga_submissions WHERE id=?",
            (submission_id,),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail=SUBMISSION_NOT_FOUND)

    sd, ed = _competencia_date_range(row["competencia"])
    year, month = sd.year, sd.month

    with tempfile.TemporaryDirectory() as tmpdir:
        ga_path = os.path.join(tmpdir, row["ga_filename"])
        async with aiofiles.open(ga_path, "wb") as f:
            await f.write(row["ga_data"])

        if is_oracle_configured():  # pragma: no cover
            person, error = await _run_oracle(ga_path, row["user_email"], sd, ed, year, month)
            if error:
                raise HTTPException(status_code=422, detail=error)
        else:
            try:
                persons, errors = _run_file(ga_path, row["competencia"])
                if not persons:
                    raise HTTPException(
                        status_code=422,
                        detail="\n".join(errors) or "Nenhum resultado",
                    )
                person = persons[0]
            except HTTPException:
                raise
            except Exception as exc:  # pragma: no cover
                raise HTTPException(status_code=422, detail=str(exc)) from exc  # pragma: no cover

    new_horas = person.get("task_total", 0.0) if person else 0.0
    new_snapshot = json.dumps(person, ensure_ascii=False)
    with _db() as conn:
        conn.execute(
            "UPDATE ga_submissions SET horas_task=?, validation_snapshot=?, status='pending', "
            "rejection_reason=NULL, submitted_at=datetime('now'),"
            " reviewed_at=NULL, reviewed_by=NULL "
            "WHERE id=?",
            (new_horas, new_snapshot, submission_id),
        )
    return {"ok": True, "horas_task": new_horas}


@app.get(
    "/api/admin/submissions/{submission_id}/preview",
    responses={
        404: {"description": SUBMISSION_NOT_FOUND},
        422: {"description": "Erro ao reprocessar"},
    },
)
async def admin_preview_submission(submission_id: int, _: AdminUser):
    with _db() as conn:
        row = conn.execute(
            "SELECT user_email, competencia, ga_filename, ga_data, validation_snapshot, "
            "approved_snapshot, status, reviewed_at, reviewed_by "
            "FROM ga_submissions WHERE id=?",
            (submission_id,),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail=SUBMISSION_NOT_FOUND)

    # Approved submissions return the immutable approved snapshot — no reprocessing.
    if row["status"] == "approved" and row["approved_snapshot"]:
        return {
            "person": json.loads(row["approved_snapshot"]),
            "submission_snapshot": json.loads(row["validation_snapshot"]),
            "is_immutable": True,
            "reviewed_at": row["reviewed_at"],
            "reviewed_by": row["reviewed_by"],
        }

    sd, ed = _competencia_date_range(row["competencia"])
    year, month = sd.year, sd.month

    with tempfile.TemporaryDirectory() as tmpdir:
        ga_path = os.path.join(tmpdir, row["ga_filename"])
        async with aiofiles.open(ga_path, "wb") as f:
            await f.write(row["ga_data"])

        if is_oracle_configured():  # pragma: no cover
            person, error = await _run_oracle(ga_path, row["user_email"], sd, ed, year, month)
            if error:
                raise HTTPException(status_code=422, detail=error)
        else:
            try:
                persons, errors = _run_file(ga_path, row["competencia"])
                if not persons:
                    raise HTTPException(
                        status_code=422,
                        detail="\n".join(errors) or "Nenhum resultado",
                    )
                person = persons[0]
            except HTTPException:
                raise
            except Exception as exc:  # pragma: no cover
                raise HTTPException(status_code=422, detail=str(exc)) from exc  # pragma: no cover

    return {
        "person": person,
        "submission_snapshot": json.loads(row["validation_snapshot"]),
        "is_immutable": False,
        "reviewed_at": row["reviewed_at"],
        "reviewed_by": row["reviewed_by"],
    }


# ── Routes: billing export (admin) ────────────────────────────────────────────

_BILLING_HEADERS = [
    "CNPJ EMPRESA",
    "Nº NOTA FISCAL",
    "CPF PROFISSIONAL",
    "Nº SUBCONTRATO",
    "VALOR HORA",
    "QUANTIDADE DE HORAS (DECIMAL)",
    "TOTAL",
    "COMPETÊNCIA EVENTO",
    "NOME DO PROFISSIONAL",
]
_CURRENCY_FMT = '"R$ "#,##0.00'
_HOURS_FMT = "#,##0.00"


def _format_cnpj(cnpj: str) -> str:
    """Format a raw CNPJ string as XX.XXX.XXX/XXXX-XX if it has 14 digits."""
    digits = _re.sub(r"\D", "", cnpj)
    if len(digits) != 14:
        return cnpj
    return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"


def _cell_display_len(cell) -> int:
    """Estimate display width of a cell value for column auto-fit."""
    if cell.value is None:
        return 0
    if isinstance(cell.value, (int, float)):
        fmt = cell.number_format or ""
        if '"R$ "' in fmt:
            return len(f"R$ {cell.value:,.2f}") + 1
        return len(f"{cell.value:,.2f}")
    return len(str(cell.value))


def _fill_billing_worksheet(ws, rows, cnpj: str, comp_label: str, total_horas: float) -> int:
    """Populate the worksheet with headers and data rows. Returns last_data_row."""
    ws.cell(row=1, column=6, value=total_horas).number_format = _HOURS_FMT

    for col, header in enumerate(_BILLING_HEADERS, start=1):
        ws.cell(row=2, column=col, value=header)

    for i, row in enumerate(rows, start=3):
        total = round(row["horas_ga"] * row["valor_hora"], 2)
        ws.cell(row=i, column=1, value=cnpj)
        ws.cell(row=i, column=2, value="")
        ws.cell(row=i, column=3, value=row["cpf"])
        ws.cell(row=i, column=4, value=row["subcontrato"])
        ws.cell(row=i, column=5, value=row["valor_hora"]).number_format = _CURRENCY_FMT
        ws.cell(row=i, column=6, value=row["horas_ga"]).number_format = _HOURS_FMT
        ws.cell(row=i, column=7, value=total).number_format = _CURRENCY_FMT
        ws.cell(row=i, column=8, value=comp_label)
        ws.cell(row=i, column=9, value=row["nome_profissional"])

    return len(rows) + 2


def _apply_billing_formatting(ws, last_data_row: int, total_nf: float, n_cols: int) -> None:
    """Add summary row, center alignment, Excel table, auto-fit columns, hide gridlines."""
    summary_row = last_data_row + 1

    ws.cell(row=summary_row, column=6, value="TOTAL NF").font = Font(bold=True)
    cell_total = ws.cell(row=summary_row, column=7, value=total_nf)
    cell_total.number_format = _CURRENCY_FMT
    cell_total.font = Font(bold=True)

    center = Alignment(horizontal="center")
    centered_cols = set(range(2, n_cols))  # cols 2-8, skipping col 1 and col 9
    for row_idx in range(1, summary_row + 1):
        for col_idx in centered_cols:
            ws.cell(row=row_idx, column=col_idx).alignment = center

    table_ref = f"A2:{get_column_letter(n_cols)}{last_data_row}"
    tab = Table(displayName="TabelaFaturamento", ref=table_ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)

    for col_cells in ws.iter_cols(min_row=1, max_row=summary_row, min_col=1, max_col=n_cols):
        max_len = max((_cell_display_len(c) for c in col_cells), default=0)
        if max_len > 0:
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = max_len + 3

    ws.sheet_view.showGridLines = False


@app.get(
    "/api/admin/submissions/{submission_id}/ga-file",
    responses={404: {"description": SUBMISSION_NOT_FOUND}},
)
async def download_submission_ga_file(submission_id: int, _: AdminUser):
    with _db() as conn:
        row = conn.execute(
            "SELECT ga_filename, ga_data FROM ga_submissions WHERE id=? AND status='approved'",
            (submission_id,),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail=SUBMISSION_NOT_FOUND)
    content_type = (
        XLSX_MEDIA_TYPE if row["ga_filename"].endswith(".xlsx") else "application/vnd.ms-excel"
    )
    return Response(
        content=row["ga_data"],
        media_type=content_type,
        headers={"Content-Disposition": f'attachment; filename="{row["ga_filename"]}"'},
    )


@app.get(
    "/api/admin/billing/{competencia}/export",
    responses={404: {"description": "Nenhuma submission aprovada para esta competência"}},
)
async def export_billing(competencia: str, _: AdminUser):
    with _db() as conn:
        cnpj_row = conn.execute("SELECT value FROM settings WHERE key='billing_cnpj'").fetchone()
        cnpj = _format_cnpj(cnpj_row["value"] if cnpj_row else "")
        rows = conn.execute(
            """
            SELECT
                s.nome_profissional,
                COALESCE(
                    json_extract(s.approved_snapshot,    '$.ga_total'),
                    json_extract(s.validation_snapshot,  '$.ga_total'),
                    json_extract(s.approved_snapshot,    '$.task_total'),
                    json_extract(s.validation_snapshot,  '$.task_total'),
                    s.horas_task
                ) AS horas_ga,
                s.competencia,
                tm.cpf,
                tm.subcontrato,
                br.valor_hora
            FROM ga_submissions s
            JOIN team_members tm ON tm.user_email = s.user_email
            JOIN billing_rates br ON br.profile   = tm.profile
            WHERE s.competencia=? AND s.status='approved'
            ORDER BY tm.subcontrato ASC
        """,
            (competencia,),
        ).fetchall()

    if not rows:
        raise HTTPException(
            status_code=404, detail="Nenhuma submission aprovada para esta competência"
        )

    comp_label = _competencia_label(competencia)
    total_horas = sum((row["horas_ga"] or 0) for row in rows)
    total_nf = round(sum((row["horas_ga"] or 0) * row["valor_hora"] for row in rows), 2)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Faturamento"

    last_data_row = _fill_billing_worksheet(ws, rows, cnpj, comp_label, total_horas)
    _apply_billing_formatting(ws, last_data_row, total_nf, len(_BILLING_HEADERS))

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = competencia.replace("-", "_")
    return Response(
        content=buf.read(),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="faturamento_{safe}.xlsx"'},
    )


# ── API key auth (machine-to-machine) ────────────────────────────────────────


def _require_api_key(request: Request) -> None:
    key = request.headers.get("X-Api-Key", "")
    expected = os.getenv("MAG_EVAL_API_KEY", "")
    if not expected or key != expected:
        raise HTTPException(status_code=401, detail="API key inválida ou ausente")


# ── Routes: client evaluations (admin) ───────────────────────────────────────

CE_NOT_FOUND = "Avaliação não encontrada"


class IngestRequest(BaseModel):
    email_body: str


class UpdateEvaluationRequest(BaseModel):
    user_email: str | None = None
    comportamental_nota: int | None = None
    comportamental_comentario: str | None = None
    tecnica_nota: int | None = None
    tecnica_comentario: str | None = None
    prazo_nota: int | None = None
    prazo_comentario: str | None = None
    comentario_geral: str | None = None


class MagEvalRequest(BaseModel):
    content: str
    subject: str
    internetMessageId: str  # noqa: N815
    receivedDateTime: str  # noqa: N815


def _ce_row_to_dict(row) -> dict:
    return dict(row) if row else {}


def _active_team_members() -> list:
    return client_evaluation_repository.list_team_members_active()


@app.post(
    "/api/admin/client-evaluations/ingest",
    status_code=201,
    responses={
        409: {"description": "Registro duplicado"},
        422: {"description": "Falha no parse do e-mail"},
    },
)
async def admin_ingest_evaluation(
    body: IngestRequest,
    admin: AdminUser,
    confirm_override: bool = False,
):
    try:
        parsed = client_evaluation_engine.parse_evaluation(body.email_body)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    active_members = _active_team_members()
    user_email, match_status = client_evaluation_engine.auto_link(
        parsed["nome_extraido"], active_members
    )

    # Deduplication rule 2 — by (user_email, competencia) when user_email resolved
    if user_email:
        existing = client_evaluation_repository.find_by_email_competencia(
            user_email, parsed["competencia"]
        )
        if existing:
            if existing["status"] == "aprovado" and confirm_override:
                pass  # allow new rascunho alongside the approved record
            elif existing["status"] == "aprovado":
                raise HTTPException(
                    status_code=409,
                    detail={
                        "message": (
                            f"Registro duplicado: {parsed['nome_extraido']} já possui"
                            f" avaliação aprovada para {parsed['competencia']}"
                        ),
                        "requires_override": True,
                    },
                )
            else:
                raise HTTPException(
                    status_code=409,
                    detail=(
                        f"Registro duplicado: {parsed['nome_extraido']} já possui"
                        f" avaliação para {parsed['competencia']}"
                    ),
                )
    else:
        # Deduplication rule 3 — by normalized nome_extraido when user_email not resolved
        if client_evaluation_repository.find_by_name_competencia(
            parsed["nome_extraido"], parsed["competencia"]
        ):
            raise HTTPException(
                status_code=409,
                detail=(
                    f"Registro duplicado: {parsed['nome_extraido']} já possui"
                    f" avaliação para {parsed['competencia']}"
                ),
            )

    record_id = client_evaluation_repository.insert_draft(
        parsed=parsed,
        user_email=user_email,
        match_status=match_status,
        raw_email_text=body.email_body,
        email_message_id=None,
        origem="manual",
        created_by=admin,
    )
    return {
        "id": record_id,
        "status": "rascunho",
        "origem": "manual",
        "competencia": parsed["competencia"],
    }


@app.get("/api/admin/client-evaluations")
async def admin_list_evaluations(
    _: AdminUser,
    status: str | None = None,
):
    rows = client_evaluation_repository.list_all(status_filter=status)
    return {"evaluations": [_ce_row_to_dict(r) for r in rows]}


# IMPORTANT: this route is registered BEFORE /{id} to avoid FastAPI path conflict.
@app.get("/api/admin/client-evaluations/dashboard")
async def admin_dashboard(
    _: AdminUser,
    user_email: str | None = None,
    competencia_from: str | None = None,
    competencia_to: str | None = None,
):
    rows = client_evaluation_repository.get_dashboard(
        user_email_filter=user_email,
        competencia_from=competencia_from,
        competencia_to=competencia_to,
    )
    return {"data": [_ce_row_to_dict(r) for r in rows]}


@app.get(
    "/api/admin/client-evaluations/{record_id}",
    responses={404: {"description": CE_NOT_FOUND}},
)
async def admin_get_evaluation(record_id: int, _: AdminUser):
    row = client_evaluation_repository.get_by_id(record_id)
    if not row:
        raise HTTPException(status_code=404, detail=CE_NOT_FOUND)
    return _ce_row_to_dict(row)


@app.patch(
    "/api/admin/client-evaluations/{record_id}",
    responses={
        404: {"description": CE_NOT_FOUND},
        409: {"description": "Registro aprovado é imutável"},
    },
)
async def admin_update_evaluation(
    record_id: int,
    body: UpdateEvaluationRequest,
    admin: AdminUser,
):
    row = client_evaluation_repository.get_by_id(record_id)
    if not row:
        raise HTTPException(status_code=404, detail=CE_NOT_FOUND)
    if row["status"] == "aprovado":
        raise HTTPException(status_code=409, detail="Registro aprovado é imutável")

    fields = body.model_dump(exclude_unset=True)
    edited_at = datetime.now(tz=UTC).isoformat()
    client_evaluation_repository.update_draft(record_id, fields, admin, edited_at)
    return {"ok": True}


@app.post(
    "/api/admin/client-evaluations/{record_id}/approve",
    responses={
        404: {"description": CE_NOT_FOUND},
        409: {"description": "Registro já aprovado"},
        422: {"description": "Profissional não vinculado"},
    },
)
async def admin_approve_evaluation(record_id: int, admin: AdminUser):
    row = client_evaluation_repository.get_by_id(record_id)
    if not row:
        raise HTTPException(status_code=404, detail=CE_NOT_FOUND)
    if row["status"] == "aprovado":
        raise HTTPException(status_code=409, detail="Registro já aprovado")
    if not row["user_email"]:
        raise HTTPException(
            status_code=422,
            detail="Profissional não vinculado — resolva o vínculo antes de aprovar",
        )
    reviewed_at = datetime.now(tz=UTC).isoformat()
    client_evaluation_repository.update_status(record_id, "aprovado", admin, reviewed_at)
    return {"ok": True}


@app.post(
    "/api/admin/client-evaluations/{record_id}/reject",
    responses={404: {"description": CE_NOT_FOUND}},
)
async def admin_reject_evaluation(record_id: int, admin: AdminUser):
    row = client_evaluation_repository.get_by_id(record_id)
    if not row:
        raise HTTPException(status_code=404, detail=CE_NOT_FOUND)
    reviewed_at = datetime.now(tz=UTC).isoformat()
    client_evaluation_repository.update_status(record_id, "rejeitado", admin, reviewed_at)
    return {"ok": True}


@app.post("/api/admin/client-evaluations/reprocess-unlinked")
async def admin_reprocess_unlinked_evaluations(admin: AdminUser):
    """Re-run auto_link() for every 'nao_vinculado' draft against the current team roster —
    picks up matches that only became possible after a collaborator was registered, had their
    nome_completo filled in, or was reactivated."""
    active_members = _active_team_members()
    edited_at = datetime.now(tz=UTC).isoformat()
    relinked = 0
    for row in client_evaluation_repository.list_unlinked_drafts():
        user_email, _match_status = client_evaluation_engine.auto_link(
            row["nome_extraido"], active_members
        )
        if user_email:
            client_evaluation_repository.relink_auto(row["id"], user_email, admin, edited_at)
            relinked += 1
    return {"ok": True, "relinked": relinked}


@app.post(
    "/api/integrations/mag-evaluation",
    status_code=201,
    responses={
        200: {"description": "Skipped (duplicate)"},
        201: {"description": "Draft created"},
        401: {"description": "API key inválida ou ausente"},
        422: {"description": "Falha no parse"},
    },
)
async def integration_mag_evaluation(
    request: Request,
    body: MagEvalRequest,
    _: Annotated[None, Depends(_require_api_key)],
):
    _skipped = JSONResponse(content={"ok": True, "skipped": True}, status_code=200)

    # Deduplication rule 1 — by internetMessageId (before parsing)
    existing_by_msg = client_evaluation_repository.find_by_message_id(body.internetMessageId)
    if existing_by_msg:
        return _skipped

    content_plain = client_evaluation_engine.strip_html(body.content)

    try:
        parsed = client_evaluation_engine.parse_evaluation(content_plain)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    active_members = _active_team_members()
    user_email, match_status = client_evaluation_engine.auto_link(
        parsed["nome_extraido"], active_members
    )

    # Deduplication rules 2 & 3 — silent skip on automated path
    if user_email:
        competencia = parsed["competencia"]
        if client_evaluation_repository.find_by_email_competencia(user_email, competencia):
            return _skipped
    else:
        if client_evaluation_repository.find_by_name_competencia(
            parsed["nome_extraido"], parsed["competencia"]
        ):
            return _skipped

    record_id = client_evaluation_repository.insert_draft(
        parsed=parsed,
        user_email=user_email,
        match_status=match_status,
        raw_email_text=content_plain,
        email_message_id=body.internetMessageId,
        origem="automacao",
        created_by="automacao",
    )
    return {"ok": True, "id": record_id}


# ── Routes: professional view ─────────────────────────────────────────────────


@app.get("/api/my-evaluations")
async def my_evaluations(email: CurrentUser):
    rows = client_evaluation_repository.get_my_evaluations(email)
    return {"evaluations": [_ce_row_to_dict(r) for r in rows]}


# ── Routes: admin — squads (tl-evaluation domain) ────────────────────────────


class CreateSquadRequest(BaseModel):
    nome: str
    tech_lead_user_email: str


class UpdateSquadRequest(BaseModel):
    nome: str | None = None
    tech_lead_user_email: str | None = None


def _validate_tech_lead_email(email: str) -> None:
    """Raise 422 if the email is not registered in team_members or users."""
    with _db() as conn:
        row = conn.execute(
            "SELECT 1 FROM users WHERE email = ?"
            " UNION ALL SELECT 1 FROM team_members WHERE user_email = ?",
            (email, email),
        ).fetchone()
    if not row:
        raise HTTPException(
            status_code=422,
            detail="tech_lead_user_email não corresponde a nenhum usuário cadastrado",
        )


@app.post(
    "/api/admin/squads",
    status_code=201,
    responses={
        422: {"description": "tech_lead_user_email não encontrado"},
        409: {"description": "Nome ou TL já cadastrado"},
    },
)
async def admin_create_squad(body: CreateSquadRequest, _: AdminUser):
    _validate_tech_lead_email(body.tech_lead_user_email)
    try:
        row = squad_repository.create_squad(body.nome, body.tech_lead_user_email)
    except sqlite3.IntegrityError as exc:
        raise HTTPException(status_code=409, detail="Nome de squad ou TL já cadastrado") from exc
    return dict(row)


@app.get("/api/admin/squads")
async def admin_list_squads(_: AdminUser):
    rows = squad_repository.list_squads()
    return {"squads": [dict(r) for r in rows]}


@app.put(
    "/api/admin/squads/{squad_id}",
    responses={
        404: {"description": SQUAD_NOT_FOUND},
        409: {"description": "Nome de squad ou TL já cadastrado"},
        422: {"description": "tech_lead_user_email não encontrado"},
    },
)
async def admin_update_squad(squad_id: int, body: UpdateSquadRequest, _: AdminUser):
    with _db() as conn:
        existing = conn.execute("SELECT 1 FROM squads WHERE id = ?", (squad_id,)).fetchone()
    if not existing:
        raise HTTPException(status_code=404, detail=SQUAD_NOT_FOUND)
    fields = body.model_dump(exclude_unset=True)
    if "tech_lead_user_email" in fields:
        _validate_tech_lead_email(fields["tech_lead_user_email"])
    try:
        if fields:
            squad_repository.update_squad(squad_id, fields)
    except sqlite3.IntegrityError as exc:
        raise HTTPException(status_code=409, detail="Nome de squad ou TL já cadastrado") from exc
    with _db() as conn:
        row = conn.execute("SELECT * FROM squads WHERE id = ?", (squad_id,)).fetchone()
    return dict(row)


@app.get(
    "/api/admin/squads/{squad_id}/members",
    responses={
        404: {"description": SQUAD_NOT_FOUND},
    },
)
async def admin_list_squad_members(squad_id: int, _: AdminUser):
    with _db() as conn:
        squad = conn.execute("SELECT 1 FROM squads WHERE id = ?", (squad_id,)).fetchone()
    if not squad:
        raise HTTPException(status_code=404, detail=SQUAD_NOT_FOUND)
    rows = squad_repository.list_squad_members_admin(squad_id)
    return {"members": [dict(r) for r in rows]}


# ── Routes: TL squads and members (tl-evaluation domain) ─────────────────────


class CreateMemberRequest(BaseModel):
    nome: str
    papel: str
    senioridade: str
    fornecedor: str


class UpdateMemberRequest(BaseModel):
    nome: str | None = None
    papel: str | None = None
    senioridade: str | None = None
    fornecedor: str | None = None
    ativo: bool | None = None


class ChatRequest(BaseModel):
    squad_id: int
    competencia: str
    squad_member_id: int
    messages: list[dict]


class ConfirmEvaluationRequest(BaseModel):
    squad_id: int
    competencia: str
    squad_member_id: int
    comportamental_nota: int
    comportamental_paragrafo: str
    tecnica_nota: int
    tecnica_paragrafo: str
    prazo_nota: int
    prazo_paragrafo: str
    janela_reduzida: bool


def _is_admin(email: str) -> bool:
    with _db() as conn:
        row = conn.execute(_SELECT_IS_ADMIN_SQL, (email,)).fetchone()
    return bool(row and row["is_admin"])


def _require_tl_or_admin(email: str, squad_id: int) -> None:
    if not (_is_admin(email) or squad_repository.is_squad_tl(email, squad_id)):
        raise HTTPException(status_code=403)


@app.get("/api/tl/squads")
async def tl_list_my_squads(email: CurrentUser):
    rows = squad_repository.list_tl_squads(email)
    return {"squads": [dict(r) for r in rows]}


@app.post(
    "/api/tl/squads/{squad_id}/members",
    status_code=201,
    responses={403: {"description": "Acesso negado"}},
)
async def tl_create_member(squad_id: int, body: CreateMemberRequest, email: CurrentUser):
    _require_tl_or_admin(email, squad_id)
    row = squad_repository.create_member(
        squad_id, body.nome, body.papel, body.senioridade, body.fornecedor
    )
    return dict(row)


@app.get(
    "/api/tl/squads/{squad_id}/members",
    responses={403: {"description": "Acesso negado"}},
)
async def tl_list_members(squad_id: int, email: CurrentUser, include_inactive: bool = False):
    _require_tl_or_admin(email, squad_id)
    rows = (
        squad_repository.list_squad_members_admin(squad_id)
        if include_inactive
        else squad_repository.list_members(squad_id)
    )
    return {"members": [dict(r) for r in rows]}


@app.patch(
    "/api/tl/squads/{squad_id}/members/{member_id}",
    responses={
        403: {"description": "Acesso negado"},
        404: {"description": MEMBER_NOT_FOUND},
    },
)
async def tl_update_member(
    squad_id: int, member_id: int, body: UpdateMemberRequest, email: CurrentUser
):
    _require_tl_or_admin(email, squad_id)
    fields = body.model_dump(exclude_unset=True)
    if "ativo" in fields:
        fields["ativo"] = int(fields["ativo"])
    updated = squad_repository.update_member(member_id, fields)
    if updated == 0:
        raise HTTPException(status_code=404, detail=MEMBER_NOT_FOUND)
    with _db() as conn:
        row = conn.execute("SELECT * FROM squad_members WHERE id = ?", (member_id,)).fetchone()
    return dict(row)


@app.get(
    "/api/tl/squads/{squad_id}/historico",
    responses={403: {"description": "Acesso negado"}},
)
async def tl_get_historico(squad_id: int, email: CurrentUser):
    _require_tl_or_admin(email, squad_id)
    rows = tl_evaluation_repository.get_historico(squad_id)
    return {"historico": [dict(r) for r in rows]}


# ── Routes: TL evaluations (tl-evaluation domain) ────────────────────────────


@app.post(
    "/api/tl-evaluations/chat",
    responses={
        403: {"description": "Acesso negado"},
        502: {"description": "Resposta inválida da IA de avaliação"},
        503: {"description": "Falha ao chamar a IA de avaliação"},
    },
)
async def tl_chat(body: ChatRequest, email: CurrentUser):
    _require_tl_or_admin(email, body.squad_id)
    request_body = tl_evaluation_engine.build_chat_request(body.messages)
    try:
        response = anthropic_client.complete(request_body)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    try:
        return tl_evaluation_engine.parse_tool_response(response)
    except ValueError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.post(
    "/api/tl-evaluations",
    status_code=201,
    responses={403: {"description": "Acesso negado"}},
)
async def tl_confirm_evaluation(body: ConfirmEvaluationRequest, email: CurrentUser):
    _require_tl_or_admin(email, body.squad_id)
    tl_evaluation_repository.upsert_evaluation(
        squad_id=body.squad_id,
        squad_member_id=body.squad_member_id,
        competencia=body.competencia,
        comportamental_nota=body.comportamental_nota,
        comportamental_paragrafo=body.comportamental_paragrafo,
        tecnica_nota=body.tecnica_nota,
        tecnica_paragrafo=body.tecnica_paragrafo,
        prazo_nota=body.prazo_nota,
        prazo_paragrafo=body.prazo_paragrafo,
        janela_reduzida=body.janela_reduzida,
        confirmed_by=email,
    )
    return {"ok": True}


@app.get(
    "/api/tl-evaluations",
    responses={403: {"description": "Acesso negado"}},
)
async def tl_list_evaluations(squad_id: int, competencia: str, email: CurrentUser):
    _require_tl_or_admin(email, squad_id)
    rows = tl_evaluation_repository.list_evaluations(squad_id, competencia)
    return {"evaluations": [dict(r) for r in rows]}


@app.post(
    "/api/tl-evaluations/{squad_id}/{competencia}/{squad_member_id}/reopen",
    responses={403: {"description": "Acesso negado"}},
)
async def tl_reopen_evaluation(
    squad_id: int, competencia: str, squad_member_id: int, email: CurrentUser
):
    _require_tl_or_admin(email, squad_id)
    tl_evaluation_repository.reopen_evaluation(squad_id, competencia, squad_member_id)
    return {"ok": True}


@app.post(
    "/api/tl-evaluations/{squad_id}/{competencia}/close",
    status_code=201,
    responses={403: {"description": "Acesso negado"}},
)
async def tl_close_competencia(squad_id: int, competencia: str, email: CurrentUser):
    _require_tl_or_admin(email, squad_id)
    try:
        row = tl_evaluation_repository.close_competencia(squad_id, competencia, email)
    except sqlite3.IntegrityError:
        row = tl_evaluation_repository.get_closure(squad_id, competencia)
    return dict(row)


@app.get(
    "/api/tl-evaluations/{squad_id}/{competencia}/close",
    responses={403: {"description": "Acesso negado"}},
)
async def tl_get_closure_status(squad_id: int, competencia: str, email: CurrentUser):
    _require_tl_or_admin(email, squad_id)
    row = tl_evaluation_repository.get_closure(squad_id, competencia)
    if not row:
        return {"closed": False}
    return {"closed": True, "closed_at": row["closed_at"], "closed_by": row["closed_by"]}


@app.delete(
    "/api/tl-evaluations/{squad_id}/{competencia}/close",
    responses={403: {"description": "Acesso negado"}},
)
async def tl_reopen_competencia(squad_id: int, competencia: str, email: CurrentUser):
    _require_tl_or_admin(email, squad_id)
    tl_evaluation_repository.reopen_competencia(squad_id, competencia)
    return {"closed": False}


@app.get(
    "/api/tl-evaluations/{squad_id}/{competencia}/pdf",
    responses={
        403: {"description": "Acesso negado"},
        404: {"description": "Competência não encerrada"},
    },
)
async def tl_get_pdf(squad_id: int, competencia: str, email: CurrentUser):
    _require_tl_or_admin(email, squad_id)
    closure = tl_evaluation_repository.get_closure(squad_id, competencia)
    if not closure:
        raise HTTPException(
            status_code=404,
            detail="Competência não encerrada — encerre antes de gerar o PDF",
        )
    with _db() as conn:
        squad = conn.execute("SELECT nome FROM squads WHERE id = ?", (squad_id,)).fetchone()
    evaluations = tl_evaluation_repository.list_evaluations(squad_id, competencia)
    pdf_bytes = tl_evaluation_pdf.generate_squad_pdf(
        squad_name=squad["nome"] if squad else str(squad_id),
        competencia=competencia,
        evaluations=[dict(e) for e in evaluations],
    )
    safe_comp = competencia.replace("-", "_")
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="avaliacao_{safe_comp}.pdf"'},
    )


# ── Startup + static ──────────────────────────────────────────────────────────


@app.on_event("startup")
def startup() -> None:
    _init_db()
    TASKS_DIR.mkdir(exist_ok=True)
    try_init_oracle_client()


if STATIC_DIR.exists():  # pragma: no cover
    if STATIC_DIR == _FRONTEND_DIST:
        # React SPA: serve /assets/* as static files; every other GET path serves
        # index.html so React Router handles client-side navigation (/results, etc.)
        assets_dir = STATIC_DIR / "assets"
        if assets_dir.exists():
            app.mount("/assets", StaticFiles(directory=assets_dir), name="assets")

        # Serve root-level static files (e.g. favicon, logo) before the SPA fallback
        for _static_file in STATIC_DIR.iterdir():
            if _static_file.is_file() and _static_file.name != "index.html":
                _name = _static_file.name

                def _make_route(static_path: Path):
                    route = "/" + static_path.name

                    @app.get(route, include_in_schema=False)
                    async def _static_root():
                        return FileResponse(static_path)

                _make_route(_static_file)

        @app.get("/{full_path:path}")
        async def spa_fallback(full_path: str):
            return FileResponse(STATIC_DIR / "index.html")

    else:
        # Legacy monolith — html=True makes / serve index.html automatically
        app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")
